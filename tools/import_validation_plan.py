#!/usr/bin/env python3
"""Read the supplied OSVP workbook without modifying it; preserve merged-cell origins."""
import argparse
import hashlib
import json
import re
from pathlib import Path
import openpyxl


def extract(workbook_path):
    book = openpyxl.load_workbook(workbook_path, data_only=False, read_only=False)
    sheet = book['Batch1']
    fields = {'I':'method','J':'thematic_target','K':'sample_requirement',
              'L':'spatial_coverage','M':'positional_target','N':'clarifications'}
    products = {}
    for row in range(1,sheet.max_row+1):
        product = sheet[f'C{row}'].value
        if not isinstance(product,str) or not re.fullmatch(r'S5-02-\d{2}',product):
            continue
        record = {'product':product,'row':row,'name':sheet[f'D{row}'].value,
                  'production_frequency_raw':sheet[f'F{row}'].value,'fields':{}}
        for column,field in fields.items():
            requested = f'{column}{row}'
            source = requested
            for merged in sheet.merged_cells.ranges:
                if requested in merged:
                    source = sheet.cell(merged.min_row,merged.min_col).coordinate
                    break
            record['fields'][field] = {'value':sheet[source].value,'cell':source,
                                      'requested_cell':requested,'sheet':sheet.title}
        products[product] = record
    book.close()
    if set(products) != {f'S5-02-{i:02}' for i in range(1,12)}:
        raise ValueError('Expected exactly the eleven S5-02 product rows')
    return {'source':workbook_path.name,
            'sha256':hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
            'sheet':'Batch1','products':products}


def apply(plan,config):
    for product,record in plan['products'].items():
        cfg = config['products'][product]
        fields = record['fields']
        target = fields['thematic_target']['value']
        match = re.fullmatch(r'(Relative RMSE|F1 score)\s*(<=|>=)\s*(\d+)%',target.strip())
        if not match:
            raise ValueError(f'Unrecognized thematic target for {product}')
        metric,op,value = match.groups()
        cfg['thematic_target'] = {'metric':'f1' if metric=='F1 score' else 'relative_rmse_pct',
            'operator':op,'value':int(value)/100 if metric=='F1 score' else int(value)}
        pos = re.fullmatch(r'(>=|>)\s*(\d+)%',fields['positional_target']['value'].strip())
        if not pos:
            raise ValueError(f'Unrecognized positional target for {product}')
        cfg['positional_target'].update(operator=pos[1],percent=int(pos[2]))
        sample = fields['sample_requirement']['value']
        cfg['sample_target'] = 3 if 'fino a 3 mensilità' in sample else int(re.match(r'\d+',sample)[0])
        cfg['per_area_target'] = 5 if '5 per ciascuna sub-AOI' in sample else None
        cfg['plan'] = record
        cfg['evaluation_scope'] = 'single_overall_product_evaluation' if 'unica valutazione complessiva' in sample else None
        cfg['kpi_aggregation'] = cfg.get('kpi_aggregation')
        cfg['required_metrics'] = [cfg['thematic_target']['metric'],'positional_accuracy_pct']
        cfg['diagnostic_metrics'] = ['precision','recall_pod','accuracy','iou','false_alarm_ratio'] if cfg['kind']=='binary' else ['rmse','mae','bias','r2']
        if product=='S5-02-04':
            cfg['reference']='Regional gauges: Emilia-Romagna, Lazio, Sicily'
            cfg['reference_regions']=['Emilia-Romagna','Lazio','Sicily']
    config['methodology_source'] = plan['source']+' [Batch1!I11:N21]'
    config['methodology_sha256'] = plan['sha256']
    config['campaign']['duration_days'] = 90
    return config


def main():
    base = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook',type=Path)
    parser.add_argument('--config',type=Path,default=base/'config/validation.json')
    args = parser.parse_args()
    plan = extract(args.workbook)
    config = apply(plan,json.loads(args.config.read_text()))
    # Complete parsing before writing either local configuration file.
    for path,data in [(base/'config/validation_plan_source.json',plan),(args.config,config)]:
        temporary = path.with_suffix('.tmp')
        temporary.write_text(json.dumps(data,indent=2,ensure_ascii=False)+'\n')
        temporary.replace(path)
    print('Imported 11 product requirements with merged-cell provenance; source workbook unchanged.')


if __name__=='__main__': main()
